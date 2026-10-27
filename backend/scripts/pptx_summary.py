#!/usr/bin/env python3
"""
pptx_summary.py
PowerPoint (.pptx) Auto Summary Extractor for SEIN PE-Portal.
Extracts defect report metadata, analysis conclusions, and countermeasures
from multiple PowerPoint files into structured JSON and formatted Excel (.xlsx).

// ponytail: Resilient dual-engine architecture:
// 1. Uses python-pptx and openpyxl if installed
// 2. Pure Python standard library (zipfile + xml.etree.ElementTree) fallback
//    guaranteeing 100% functionality even in minimal offline/baremetal environments.
"""

import os
import sys
import json
import argparse
import re
from datetime import datetime
import zipfile
import xml.etree.ElementTree as ET

# Try optional third-party libraries
try:
    from pptx import Presentation
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    import pandas as pd
    HAS_PANDAS = True
except ImportError:
    HAS_PANDAS = False


def extract_sections_from_text(text, data):
    """Mengekstrak conclusion dan action/countermeasure dari potongan teks."""
    if not text:
        return
    text_str = str(text).strip()
    text_lower = text_str.lower()

    # 1. Ekstraksi Conclusion
    if "conclusion" in text_lower:
        idx = text_lower.find("conclusion")
        val = text_str[idx + len("conclusion"):].strip().lstrip(":-").strip()
        if val and val not in data["Symptom/Analysis/Conclusion"]:
            data["Symptom/Analysis/Conclusion"] += ("\n" if data["Symptom/Analysis/Conclusion"] else "") + val

    # 2. Ekstraksi Action / Countermeasure
    if "action" in text_lower or "countermeasure" in text_lower:
        for kw in ["action & countermeasure", "countermeasure", "action"]:
            if kw in text_lower:
                idx = text_lower.find(kw)
                val = text_str[idx + len(kw):].strip().lstrip(":-").strip()
                if val and val not in data["Action & Countermeasure"]:
                    data["Action & Countermeasure"] += ("\n" if data["Action & Countermeasure"] else "") + val
                break


def match_table_metadata(row_texts, data):
    """Mencocokkan baris tabel metadata standar Samsung OQC / PE (Model, Problem, CN, dll)."""
    for i, cell_text in enumerate(row_texts):
        cell_lower = cell_text.lower().strip()
        if i + 1 < len(row_texts):
            next_val = row_texts[i + 1].strip()
            if not next_val:
                continue

            if cell_lower == "model" and not data["Model"]:
                data["Model"] = next_val
            elif cell_lower == "problem" and not data["Problem"]:
                data["Problem"] = next_val
            elif cell_lower == "cn" and not data["CN"]:
                data["CN"] = next_val
            elif "defect date" in cell_lower and not data["Defect Date"]:
                data["Defect Date"] = next_val
            elif "defect location" in cell_lower and not data["Defect Location"]:
                data["Defect Location"] = next_val
            elif "defect qty" in cell_lower and not data["Defect Qty"]:
                data["Defect Qty"] = next_val


def parse_with_python_pptx(file_path):
    """Mengekstrak informasi PowerPoint menggunakan pustaka python-pptx."""
    prs = Presentation(file_path)
    data = {
        "File Name": os.path.basename(file_path),
        "Model": "",
        "Problem": "",
        "CN": "",
        "Defect Date": "",
        "Defect Location": "",
        "Defect Qty": "",
        "Symptom/Analysis/Conclusion": "",
        "Action & Countermeasure": ""
    }

    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                extract_sections_from_text(shape.text_frame.text, data)
            if shape.has_table:
                for row in shape.table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells]
                    match_table_metadata(row_texts, data)
                    for cell_text in row_texts:
                        extract_sections_from_text(cell_text, data)

    return data


def parse_with_xml_fallback(file_path):
    """
    Fallback parser menggunakan standard library zipfile + xml.etree.ElementTree.
    Membaca struktur OpenXML ppt/slides/slide*.xml secara langsung tanpa dependencies.
    """
    data = {
        "File Name": os.path.basename(file_path),
        "Model": "",
        "Problem": "",
        "CN": "",
        "Defect Date": "",
        "Defect Location": "",
        "Defect Qty": "",
        "Symptom/Analysis/Conclusion": "",
        "Action & Countermeasure": ""
    }

    ns = {
        'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
        'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
    }

    with zipfile.ZipFile(file_path, 'r') as z:
        # Cari semua slide
        slide_names = [name for name in z.namelist() if re.match(r'ppt/slides/slide\d+\.xml', name)]
        # Urutkan berdasarkan nomor slide
        slide_names.sort(key=lambda s: int(re.search(r'\d+', s).group()) if re.search(r'\d+', s) else 0)

        for slide_name in slide_names:
            try:
                xml_data = z.read(slide_name)
                root = ET.fromstring(xml_data)

                # 1. Ekstrak teks dari shapes & text frames
                for sp in root.findall('.//p:sp', ns):
                    paragraphs = sp.findall('.//a:p', ns)
                    para_texts = []
                    for p in paragraphs:
                        runs = p.findall('.//a:t', ns)
                        text_val = "".join([t.text or "" for t in runs]).strip()
                        if text_val:
                            para_texts.append(text_val)
                    full_shape_text = "\n".join(para_texts)
                    extract_sections_from_text(full_shape_text, data)

                # 2. Ekstrak tabel
                for tbl in root.findall('.//a:tbl', ns):
                    for tr in tbl.findall('.//a:tr', ns):
                        row_texts = []
                        for tc in tr.findall('.//a:tc', ns):
                            tc_runs = tc.findall('.//a:t', ns)
                            tc_text = "".join([t.text or "" for t in tc_runs]).strip()
                            row_texts.append(tc_text)
                        
                        match_table_metadata(row_texts, data)
                        for cell_text in row_texts:
                            extract_sections_from_text(cell_text, data)

            except Exception:
                continue

    return data


def parse_pptx_report(file_path):
    """Mengekstrak informasi dari satu berkas PowerPoint (.pptx)."""
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    # Coba gunakan python-pptx jika tersedia
    if HAS_PPTX:
        try:
            return parse_with_python_pptx(file_path)
        except Exception as e:
            # Fallback ke pure XML jika python-pptx gagal
            try:
                return parse_with_xml_fallback(file_path)
            except Exception:
                raise e
    else:
        # Gunakan pure XML parser
        return parse_with_xml_fallback(file_path)


def save_to_formatted_excel(results, output_path):
    """Menyimpan list dictionary hasil ke Excel dan merapikan tampilannya."""
    headers = [
        "File Name",
        "Model",
        "Problem",
        "CN",
        "Defect Date",
        "Defect Location",
        "Defect Qty",
        "Symptom/Analysis/Conclusion",
        "Action & Countermeasure"
    ]

    if HAS_OPENPYXL:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Summary Report"

        # Styling Header
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Calibri", size=11, color="FFFFFF", bold=True)
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        ws.append(headers)

        for col_idx, cell in enumerate(ws[1], start=1):
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.row_dimensions[1].height = 28

        # Isi Data
        for row_idx, item in enumerate(results, start=2):
            row_vals = [item.get(h, "") for h in headers]
            ws.append(row_vals)
            ws.row_dimensions[row_idx].height = 45  # Tinggi cukup untuk multi-line text

            # Zebra stripe rows for readability
            bg_color = "F9FBFD" if row_idx % 2 == 0 else "FFFFFF"
            row_fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")

            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.fill = row_fill
                cell.border = thin_border
                cell.font = Font(name="Calibri", size=10)
                # Align left for long text, center for short metadata
                if col_idx in [2, 4, 5, 6, 7]:  # Model, CN, Date, Location, Qty
                    cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

        # Atur lebar kolom
        for col in ws.columns:
            max_len = 0
            for cell in col:
                val = str(cell.value or "")
                # ambil baris terpanjang jika ada multiline
                lines = val.split('\n')
                line_max = max(len(l) for l in lines) if lines else 0
                max_len = max(max_len, line_max)

            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = min(max(max_len + 5, 16), 55)

        # Freeze header row
        ws.freeze_panes = "A2"

        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        wb.save(output_path)

    else:
        # Fallback: Buat genuine .xlsx menggunakan zipfile + XML tanpa external dependency
        try:
            create_pure_xlsx(results, headers, output_path)
        except Exception:
            # Secondary fallback: CSV
            import csv
            csv_path = output_path.replace('.xlsx', '.csv')
            os.makedirs(os.path.dirname(os.path.abspath(csv_path)), exist_ok=True)
            with open(csv_path, 'w', newline='', encoding='utf-8-sig') as f:
                writer = csv.DictWriter(f, fieldnames=headers)
                writer.writeheader()
                for row in results:
                    writer.writerow({h: row.get(h, "") for h in headers})
            output_path = csv_path

    return output_path


def create_pure_xlsx(rows, headers, output_path):
    """Membuat file .xlsx valid secara langsung dengan zipfile dan OpenXML XML."""
    import xml.sax.saxutils as saxutils

    content_types = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
  <Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
  <Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>
</Types>'''

    rels = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
</Relationships>'''

    wb_rels = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
  <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
</Relationships>'''

    workbook = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <sheets>
    <sheet name="Summary Report" sheetId="1" r:id="rId1"/>
  </sheets>
</workbook>'''

    styles = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <fonts count="2">
    <font><name val="Calibri"/><sz val="11"/></font>
    <font><b/><color rgb="FFFFFFFF"/><name val="Calibri"/><sz val="11"/></font>
  </fonts>
  <fills count="3">
    <fill><patternFill patternType="none"/></fill>
    <fill><patternFill patternType="gray125"/></fill>
    <fill><patternFill patternType="solid"><fgColor rgb="FF1F4E78"/></patternFill></fill>
  </fills>
  <borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>
  <cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>
  <cellXfs count="3">
    <xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>
    <xf numFmtId="0" fontId="1" fillId="2" borderId="0" xfId="0" applyFont="1" applyFill="1" applyAlignment="1">
      <alignment horizontal="center" vertical="center" wrapText="1"/>
    </xf>
    <xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0" applyAlignment="1">
      <alignment horizontal="left" vertical="top" wrapText="1"/>
    </xf>
  </cellXfs>
</styleSheet>'''

    sheet_data = ['<sheetData>']
    # Header row
    sheet_data.append('<row r="1" ht="28" customHeight="1">')
    for col_i, h in enumerate(headers, 1):
        col_letter = chr(64 + col_i) if col_i <= 26 else f'A{chr(64 + col_i - 26)}'
        esc = saxutils.escape(str(h))
        sheet_data.append(f'<c r="{col_letter}1" s="1" t="inlineStr"><is><t>{esc}</t></is></c>')
    sheet_data.append('</row>')

    for row_i, row in enumerate(rows, 2):
        sheet_data.append(f'<row r="{row_i}" ht="45" customHeight="1">')
        for col_i, h in enumerate(headers, 1):
            col_letter = chr(64 + col_i) if col_i <= 26 else f'A{chr(64 + col_i - 26)}'
            val = saxutils.escape(str(row.get(h, '')))
            sheet_data.append(f'<c r="{col_letter}{row_i}" s="2" t="inlineStr"><is><t>{val}</t></is></c>')
        sheet_data.append('</row>')
    sheet_data.append('</sheetData>')

    sheet1 = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <cols>
    <col min="1" max="1" width="25" customWidth="1"/>
    <col min="2" max="2" width="18" customWidth="1"/>
    <col min="3" max="3" width="30" customWidth="1"/>
    <col min="4" max="4" width="18" customWidth="1"/>
    <col min="5" max="5" width="16" customWidth="1"/>
    <col min="6" max="6" width="20" customWidth="1"/>
    <col min="7" max="7" width="14" customWidth="1"/>
    <col min="8" max="8" width="45" customWidth="1"/>
    <col min="9" max="9" width="45" customWidth="1"/>
  </cols>
  {''.join(sheet_data)}
</worksheet>'''

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    with zipfile.ZipFile(output_path, 'w', compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml', content_types)
        z.writestr('_rels/.rels', rels)
        z.writestr('xl/workbook.xml', workbook)
        z.writestr('xl/_rels/workbook.xml.rels', wb_rels)
        z.writestr('xl/styles.xml', styles)
        z.writestr('xl/worksheets/sheet1.xml', sheet1)



def main():
    parser = argparse.ArgumentParser(description="SEIN PE-Portal PowerPoint Auto Summary Extractor")
    parser.add_argument("inputs", nargs="*", help="File path atau direktori input berkas .pptx")
    parser.add_argument("--input-dir", help="Direktori berisi kumpulan file .pptx")
    parser.add_argument("--output-excel", help="Path tujuan output berkas Excel (.xlsx)")
    parser.add_argument("--output-json", help="Path tujuan output berkas JSON hasil ekstraksi")
    parser.add_argument("--json-input", help="File JSON berisi daftar path berkas .pptx")
    args = parser.parse_args()

    files_to_process = []

    # 1. Kumpulkan file dari argumen
    if args.json_input and os.path.exists(args.json_input):
        try:
            with open(args.json_input, 'r', encoding='utf-8') as f:
                data_in = json.load(f)
                if isinstance(data_in, list):
                    files_to_process.extend(data_in)
                elif isinstance(data_in, dict) and "files" in data_in:
                    files_to_process.extend(data_in["files"])
        except Exception as e:
            print(json.dumps({"success": False, "error": f"Failed to parse json-input: {e}"}))
            sys.exit(1)

    if args.input_dir and os.path.isdir(args.input_dir):
        for root_dir, _, filenames in os.walk(args.input_dir):
            for fn in filenames:
                if fn.lower().endswith(".pptx") and not fn.startswith("~$"):
                    files_to_process.append(os.path.join(root_dir, fn))

    for item in args.inputs:
        if os.path.isfile(item) and item.lower().endswith(".pptx") and not os.path.basename(item).startswith("~$"):
            files_to_process.append(item)
        elif os.path.isdir(item):
            for root_dir, _, filenames in os.walk(item):
                for fn in filenames:
                    if fn.lower().endswith(".pptx") and not fn.startswith("~$"):
                        files_to_process.append(os.path.join(root_dir, fn))

    # Hilangkan duplikasi sambil mempertahankan urutan
    seen = set()
    unique_files = []
    for fp in files_to_process:
        norm = os.path.abspath(fp)
        if norm not in seen and os.path.exists(norm):
            seen.add(norm)
            unique_files.append(norm)

    if not unique_files:
        res = {
            "success": False,
            "total": 0,
            "succeeded": 0,
            "failed": 0,
            "error": "No valid .pptx files found to process",
            "items": [],
            "errors": []
        }
        print(json.dumps(res, indent=2))
        sys.exit(0)

    summary_results = []
    failed_files = []

    for fp in unique_files:
        fn = os.path.basename(fp)
        try:
            item_data = parse_pptx_report(fp)
            summary_results.append(item_data)
        except Exception as e:
            failed_files.append({
                "file_name": fn,
                "error": str(e)
            })

    # Tentukan path output excel jika diminta
    output_excel_path = args.output_excel
    if not output_excel_path and summary_results:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_excel_path = f"/data/files/summary_report_{timestamp}.xlsx"

    actual_excel_path = ""
    if output_excel_path and summary_results:
        try:
            actual_excel_path = save_to_formatted_excel(summary_results, output_excel_path)
        except Exception as e:
            failed_files.append({
                "file_name": "Excel_Export",
                "error": f"Failed to generate Excel file: {e}"
            })

    output_payload = {
        "success": True,
        "total": len(unique_files),
        "succeeded": len(summary_results),
        "failed": len(failed_files),
        "excel_file": os.path.basename(actual_excel_path) if actual_excel_path else "",
        "excel_path": actual_excel_path,
        "items": summary_results,
        "errors": failed_files
    }

    if args.output_json:
        try:
            with open(args.output_json, 'w', encoding='utf-8') as f:
                json.dump(output_payload, f, indent=2, ensure_ascii=False)
        except Exception:
            pass

    # Print output payload to stdout
    print(json.dumps(output_payload, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
